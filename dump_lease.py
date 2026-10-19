import json
import openpyxl

wb = openpyxl.load_workbook('华润融资租赁报表 .xlsx', data_only=True)
print('sheets:', wb.sheetnames)
out = {}
for ws in wb.worksheets:
    print('---', repr(ws.title), 'max_row:', ws.max_row, 'max_col:', ws.max_column, 'merged:', len(ws.merged_cells.ranges))
    grid = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column, values_only=True):
        grid.append([('' if v is None else str(v)) for v in row])
    # apply merges: top-left value fills range
    for mr in ws.merged_cells.ranges:
        top = grid[mr.min_row-1][mr.min_col-1]
        for r in range(mr.min_row, mr.max_row+1):
            for c in range(mr.min_col, mr.max_col+1):
                grid[r-1][c-1] = top
    out[ws.title] = grid

with open('lease_grid.json', 'w', encoding='utf-8') as f:
    json.dump(out, f, ensure_ascii=False, indent=1)
print('done -> lease_grid.json')
for k, g in out.items():
    print(k, '->', len(g), 'rows')
