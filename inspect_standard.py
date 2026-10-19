# -*- coding: utf-8 -*-
"""检查 金控指标数据梳理表-指标标准版【清洁版】.xlsx 结构"""
import openpyxl
import json
import os

BASE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(BASE, "金控指标数据梳理表-指标标准版【清洁版】 .xlsx")

wb = openpyxl.load_workbook(SRC, data_only=True)
print("sheetnames:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print("=" * 80)
    print(f"SHEET [{name}] max_row={ws.max_row} max_col={ws.max_column}")
    print("merged:", [str(r) for r in ws.merged_cells.ranges][:60])
    # 打印前 15 行
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 15), max_col=ws.max_column):
        vals = []
        for cell in row:
            v = cell.value
            if v is None:
                vals.append(None)
            elif isinstance(v, str):
                vals.append(v.strip() or None)
            elif isinstance(v, float) and v == int(v):
                vals.append(int(v))
            else:
                vals.append(v)
        if any(x is not None for x in vals):
            print(f"  R{row[0].row}:", vals)
