# -*- coding: utf-8 -*-
"""提取每张子表指标名称列的层级格式信息（加粗/缩进/前导空格）"""
import openpyxl
import json
import os

BASE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(BASE, "金控战略部银行指标体系模板.xlsx")
OUT = os.path.join(BASE, "hierarchy.json")

wb = openpyxl.load_workbook(SRC)  # 需要格式信息，不用 data_only

# 各 sheet 指标名称所在列(0基) 与首个数据行
NAME_COL = {"01": 1, "02": 1, "03": 2, "04": 3, "05": 3, "06": 3, "07": 3, "08": 3, "09": 3}
FIRST_ROW = {"01": 4, "02": 5, "03": 5, "04": 5, "05": 4, "06": 5, "07": 4, "08": 5, "09": 5}

result = {}
for sn in wb.sheetnames:
    ws = wb[sn]
    key = sn[:2]
    col_idx = NAME_COL[key] + 1
    rows = []
    for r in range(FIRST_ROW[key], ws.max_row + 1):
        cell = ws.cell(row=r, column=col_idx)
        v = cell.value
        if v is None or str(v).strip() == "":
            continue
        raw = str(v)
        lead = len(raw) - len(raw.lstrip())  # 前导空格数
        al = cell.alignment
        rows.append({
            "row": r,
            "name": raw.strip(),
            "bold": bool(cell.font.bold) if cell.font and cell.font.bold is not None else False,
            "indent": float(al.indent) if al and al.indent else 0,
            "lead": lead
        })
    result[sn] = rows

with open(OUT, "w", encoding="utf-8") as f:
    json.dump(result, f, ensure_ascii=False, indent=1)

for sn, rows in result.items():
    lv = sorted({(x["bold"], x["indent"], x["lead"]) for x in rows})
    print(sn, "格式组合:", lv)
