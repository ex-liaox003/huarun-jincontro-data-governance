import openpyxl
import json

wb = openpyxl.load_workbook('金控战略部银行指标体系模板.xlsx')

result = {}
result['sheets'] = wb.sheetnames

for name in wb.sheetnames:
    ws = wb[name]
    sheet_data = {
        'max_row': ws.max_row,
        'max_col': ws.max_column,
        'rows': []
    }
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 50), values_only=False):
        vals = [str(cell.value) if cell.value is not None else '' for cell in row]
        if any(v for v in vals):
            sheet_data['rows'].append(vals)
    result[name] = sheet_data

print(json.dumps(result, ensure_ascii=False, indent=2))
