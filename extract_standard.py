# -*- coding: utf-8 -*-
"""完整提取 金控指标数据梳理表-指标标准版【清洁版】.xlsx → indicator_standard_raw.json (UTF-8)"""
import openpyxl
import json
import os

BASE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(BASE, "金控指标数据梳理表-指标标准版【清洁版】 .xlsx")
OUT = os.path.join(BASE, "indicator_standard_raw.json")

wb = openpyxl.load_workbook(SRC, data_only=True)

result = {"file": os.path.basename(SRC), "sheet_count": len(wb.sheetnames), "sheets": []}

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    sheet_info = {
        "name": sheet_name,
        "max_row": ws.max_row,
        "max_col": ws.max_column,
        "merged_ranges": [str(r) for r in ws.merged_cells.ranges],
        "rows": []
    }
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        row_vals = []
        all_empty = True
        for cell in row:
            v = cell.value
            if v is None:
                row_vals.append(None)
            else:
                if isinstance(v, str):
                    v = v.strip()
                    if v == "":
                        row_vals.append(None)
                        continue
                elif isinstance(v, float):
                    if v == int(v):
                        v = int(v)
                row_vals.append(v)
                all_empty = False
        if not all_empty:
            sheet_info["rows"].append({"row_num": row[0].row, "cells": row_vals})
    result["sheets"].append(sheet_info)

with open(OUT, "w", encoding="utf-8") as f:
    json.dump(result, f, ensure_ascii=False, indent=1)

print("OK rows:", [len(s["rows"]) for s in result["sheets"]])
