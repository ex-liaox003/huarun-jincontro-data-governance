# -*- coding: utf-8 -*-
"""
重新完整提取 金控战略部银行指标体系模板.xlsx
输出: excel_fresh.json
包含: 每个 sheet 的准确名称、所有合并单元格信息、逐行逐格内容(含表头行与数据行)
"""
import openpyxl
import json
import sys
import os

BASE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(BASE, "金控战略部银行指标体系模板.xlsx")
OUT = os.path.join(BASE, "excel_fresh.json")

wb = openpyxl.load_workbook(SRC, data_only=True)

result = {"file": SRC, "sheet_count": len(wb.sheetnames), "sheets": []}

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    sheet_info = {
        "name": sheet_name,
        "max_row": ws.max_row,
        "max_col": ws.max_column,
        "merged_ranges": [str(r) for r in ws.merged_cells.ranges],
        "rows": []
    }
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        row_vals = []
        all_empty = True
        for cell in row:
            v = cell.value
            if v is None:
                row_vals.append(None)
            else:
                if isinstance(v, str):
                    v = v.strip()
                    if v == "":
                        row_vals.append(None)
                        continue
                elif isinstance(v, float):
                    if v == int(v):
                        v = int(v)
                row_vals.append(v)
                all_empty = False
        if not all_empty:
            sheet_info["rows"].append({"row_num": row[0].row, "cells": row_vals})
    result["sheets"].append(sheet_info)

with open(OUT, "w", encoding="utf-8") as f:
    json.dump(result, f, ensure_ascii=False, indent=1)

# 控制台摘要
print("OK")
print("sheet_count:", len(wb.sheetnames))
for s in result["sheets"]:
    print(f"- [{s['name']}] rows={len(s['rows'])} max_col={s['max_col']} merged={len(s['merged_ranges'])}")
