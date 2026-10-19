# -*- coding: utf-8 -*-
import json, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl

FN = '金控&银行指标梳理梳理表(金控风险部).xlsx'
wb = openpyxl.load_workbook(FN, data_only=True)
print("SHEETS:", wb.sheetnames)
for ws in wb.worksheets:
    print(f"--- {ws.title!r} rows={ws.max_row} cols={ws.max_column}")

ws = wb.worksheets[1]
print("=== SECOND SHEET:", ws.title)
for r in range(1, min(ws.max_row, 8) + 1):
    vals = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(r, c).value
        vals.append('' if v is None else str(v).replace('\n', '\\n'))
    print(f"ROW{r}: " + " | ".join(vals))
