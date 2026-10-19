# -*- coding: utf-8 -*-
"""分析指标标准清单字段分布 → standard_analysis.json"""
import openpyxl
import json
import os
from datetime import datetime, date

BASE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(BASE, "金控指标数据梳理表-指标标准版【清洁版】 .xlsx")
OUT = os.path.join(BASE, "standard_analysis.json")

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb[wb.sheetnames[0]]

header = []
for cell in ws[2]:
    v = cell.value
    header.append(str(v).replace("\n", "/").strip() if v is not None else None)

# 关键列索引（0基）
KEY_COLS = {
    "指标主题": 2, "一级分类": 3, "二级分类": 4, "三级分类": 5,
    "指标类型": 9, "指标口径": 10, "指标单位": 12, "业务单元": 13,
    "是否监管指标": 15, "报送监管机构": 16, "更新频次": 17, "时间维度": 18,
    "来源金控系统名称": 25, "是否入数据湖": 27, "数据类型": 28,
    "取数方式": 41, "填报部门": 43, "指标管理部门": 45, "指标状态": 46,
    "是否共享": 47, "版本号": 55, "安全级别": 56
}

def dist_of(col):
    d = {}
    for r in range(3, ws.max_row + 1):
        v = ws.cell(row=r, column=col + 1).value
        if v is None:
            continue
        if isinstance(v, str):
            v = v.strip()
            if not v:
                continue
        else:
            v = str(v)
        d[v] = d.get(v, 0) + 1
    return dict(sorted(d.items(), key=lambda kv: -kv[1]))

rows_data = []
empty_rows = []
for r in range(3, ws.max_row + 1):
    code = ws.cell(row=r, column=2).value
    name = ws.cell(row=r, column=7).value
    seq = ws.cell(row=r, column=1).value
    if code is None and name is None:
        empty_rows.append(r)
        continue
    rows_data.append({"row": r, "seq": seq, "code": code, "name": name})

analysis = {
    "sheet": ws.title,
    "max_row": ws.max_row,
    "max_col": ws.max_column,
    "header": header,
    "data_row_count": len(rows_data),
    "empty_rows": empty_rows,
    "first_row": rows_data[0] if rows_data else None,
    "last_row": rows_data[-1] if rows_data else None,
    "code_prefixes": dist_of(1),
    "distributions": {k: dist_of(ci) for k, ci in KEY_COLS.items()},
}

with open(OUT, "w", encoding="utf-8") as f:
    json.dump(analysis, f, ensure_ascii=False, indent=1)
print("OK")
