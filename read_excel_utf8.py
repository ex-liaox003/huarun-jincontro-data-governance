import openpyxl
import json
import os

base_dir = os.path.dirname(os.path.abspath(__file__))
xlsx_path = os.path.join(base_dir, '金控战略部银行指标体系模板.xlsx')
output_path = os.path.join(base_dir, 'output.json')

wb = openpyxl.load_workbook(xlsx_path)

result = {}
result['sheets'] = wb.sheetnames

for name in wb.sheetnames:
    ws = wb[name]
    sheet_data = {
        'max_row': ws.max_row,
        'max_col': ws.max_column,
        'rows': []
    }
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 50), values_only=False):
        vals = [str(cell.value) if cell.value is not None else '' for cell in row]
        if any(v for v in vals):
            sheet_data['rows'].append(vals)
    result[name] = sheet_data

with open(output_path, 'w', encoding='utf-8') as f:
    json.dump(result, f, ensure_ascii=False, indent=2)

print("Done - written to", output_path)
