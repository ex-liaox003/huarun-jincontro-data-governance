# -*- coding: utf-8 -*-
"""标记《银行指标》sheet 中的重复指标：B列=是，C列=重复名称。另存新文件。"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl

SRC = '金控&银行指标梳理梳理表(金控风险部).xlsx'
DST = '金控&银行指标梳理梳理表(金控风险部)-重复标记.xlsx'
SHEET = '银行指标'
HEADER_ROWS = 3          # 数据从第4行开始
COL_B, COL_C, COL_I = 2, 3, 9   # 是否重复 / 重复名称 / 指标建议名称

wb = openpyxl.load_workbook(SRC)
ws = wb[SHEET]

# 1. 收集名称 -> [行号]
groups = {}
for r in range(HEADER_ROWS + 1, ws.max_row + 1):
    name = str(ws.cell(r, COL_I).value or '').replace(' ', '').replace('　', '')
    if not name:
        continue
    groups.setdefault(name, []).append(r)

# 2. 重复组：同名多行
dup_groups = {n: rs for n, rs in groups.items() if len(rs) > 1}

marked, kept_alias = 0, 0
for name, rs in dup_groups.items():
    for r in rs:
        ws.cell(r, COL_B).value = '是'
        cur = str(ws.cell(r, COL_C).value or '').strip()
        if cur:                       # 已有手写别名（如 银行合计不良下迁），保留
            kept_alias += 1
        else:
            ws.cell(r, COL_C).value = name
        marked += 1

# 3. 原有"是"但不在同名组里的行（跨名称别名），保留不动
pre = set()
for r in range(HEADER_ROWS + 1, ws.max_row + 1):
    if str(ws.cell(r, COL_B).value or '').strip() == '是':
        pre.add(r)
all_group_rows = {r for rs in dup_groups.values() for r in rs}
kept_cross = sorted(pre - all_group_rows)

wb.save(DST)
print(f'重复组数: {len(dup_groups)}')
print(f'标记行数: {len(all_group_rows)}')
print(f'其中保留原别名C值: {kept_alias}')
print(f'保留的跨名称人工标记(不在同名组): {kept_cross}')
for n, rs in sorted(dup_groups.items(), key=lambda kv: kv[1][0]):
    print(f'  {n}: 行 {rs}')
