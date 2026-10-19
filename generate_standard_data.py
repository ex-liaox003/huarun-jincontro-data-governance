# -*- coding: utf-8 -*-
"""
生成 indicator_standard_data.js
来源: 金控指标数据梳理表-指标标准版【清洁版】 .xlsx (sheet: 银行指标清单)
结构: INDICATOR_STANDARD = { meta: {...}, items: [...] }
- items: 159 条指标标准，字段 key 见 FIELD_MAP
- meta: 来源信息、一级分类顺序、分组字段定义、统计信息
"""
import openpyxl
import json
import os
from datetime import datetime, date
from collections import OrderedDict

BASE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(BASE, "金控指标数据梳理表-指标标准版【清洁版】 .xlsx")
OUT = os.path.join(BASE, "indicator_standard_data.js")

# (key, 列索引0基, 分组, 展示名)
FIELD_MAP = [
    ("no",        0,  "basic",      "序号"),
    ("code",      1,  "basic",      "指标标准编号"),
    ("theme",     2,  "business",   "指标主题"),
    ("cat1",      3,  "business",   "一级分类"),
    ("cat2",      4,  "business",   "二级分类"),
    ("cat3",      5,  "business",   "三级分类"),
    ("name",      6,  "business",   "指标名称"),
    ("nameEn",    7,  "business",   "指标英文名称"),
    ("definition",8,  "business",   "指标定义"),
    ("type",      9,  "business",   "指标类型"),
    ("caliber",   10, "business",   "指标口径"),
    ("formula",   11, "business",   "指标计算公式"),
    ("unit",      12, "business",   "指标单位（金控展示）"),
    ("bizUnit",   13, "business",   "业务单元"),
    ("dataRange", 14, "business",   "数据范围"),
    ("isRegulatory",15,"business",  "是否监管指标"),
    ("regulator", 16, "business",   "报送监管机构"),
    ("freq",      17, "business",   "更新频次"),
    ("timeDim",   18, "business",   "时间维度"),
    ("bizNote",   19, "business",   "业务属性备注"),
    ("srcSystem", 20, "business",   "指标来源业务单元系统名称"),
    ("srcReport", 21, "business",   "指标来源业务单元系统报表"),
    ("supplyDept",22, "business",   "业务单元供数部门"),
    ("bizContact",23, "business",   "业务单元业务对接人"),
    ("itContact", 24, "business",   "业务单元IT对接人"),
    ("jkSystem",  25, "technical",  "来源金控系统名称"),
    ("jkReport",  26, "technical",  "来源金控报表"),
    ("inDataLake",27, "technical",  "是否入数据湖"),
    ("dataType",  28, "technical",  "数据类型"),
    ("dataFormat",29, "technical",  "数据格式"),
    ("valueDomain",30,"technical",  "值域"),
    ("storeUnit", 31, "technical",  "计量单位（数据库存储最小单位）"),
    ("lakeDb",    32, "technical",  "唯一性识别规则数据湖库名"),
    ("lakeTable", 33, "technical",  "唯一性识别规则数据湖表名"),
    ("pushTable", 34, "technical",  "推送表名"),
    ("srcCode",   35, "technical",  "指标来源编号"),
    ("companyCode",36,"technical",  "唯一性识别规则公司code"),
    ("reportCodeField",37,"technical","唯一性识别规则报表code字段"),
    ("timeDimField",38,"technical", "唯一性识别规则时间维度字段"),
    ("statTimeField",39,"technical","唯一性识别规则统计时间字段"),
    ("timeType",  40, "technical",  "时间类型"),
    ("fetchMethod",41,"technical",  "取数方式"),
    ("techNote",  42, "technical",  "技术属性备注"),
    ("reportDept",43, "management", "填报部门"),
    ("owner",     44, "management", "指标责任人"),
    ("mgmtDept",  45, "management", "指标管理部门"),
    ("status",    46, "management", "指标状态"),
    ("isShared",  47, "management", "是否共享"),
    ("creator",   48, "management", "创建人"),
    ("createDate",49, "management", "创建日期"),
    ("updater",   50, "management", "更新人"),
    ("updateDate",51, "management", "更新日期"),
    ("publishDate",52,"management", "发布日期"),
    ("disableDate",53,"management", "停用日期"),
    ("disableReason",54,"management","停用原因"),
    ("version",   55, "management", "版本号"),
    ("securityLevel",56,"management","安全级别"),
]

GROUP_LABELS = OrderedDict([
    ("basic", "基础信息"),
    ("business", "业务属性"),
    ("technical", "技术属性"),
    ("management", "管理属性"),
])


def clean(v):
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v == int(v):
        return int(v)
    if isinstance(v, str):
        v = v.strip()
        return v if v else None
    return v


wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb[wb.sheetnames[0]]

items = []
cat1_order = []
cat1_cat2 = OrderedDict()
counters = {"cat1": {}, "type": {}, "caliber": {}, "freq": {}, "unit": {}, "cat2": {}}

for r in range(3, ws.max_row + 1):
    code = clean(ws.cell(row=r, column=2).value)
    name = clean(ws.cell(row=r, column=7).value)
    if code is None and name is None:
        continue
    item = {}
    for key, ci, group, label in FIELD_MAP:
        item[key] = clean(ws.cell(row=r, column=ci + 1).value)
    items.append(item)
    c1, c2 = item.get("cat1"), item.get("cat2")
    if c1 and c1 not in cat1_order:
        cat1_order.append(c1)
        cat1_cat2[c1] = []
    if c1 and c2 and c2 not in cat1_cat2[c1]:
        cat1_cat2[c1].append(c2)
    for fld, bucket in (("cat1", "cat1"), ("type", "type"), ("caliber", "caliber"),
                        ("freq", "freq"), ("unit", "unit"), ("cat2", "cat2")):
        v = item.get(fld)
        if v:
            counters[bucket][v] = counters[bucket].get(v, 0) + 1

meta = {
    "sourceFile": "金控指标数据梳理表-指标标准版【清洁版】.xlsx",
    "sheetName": ws.title,
    "version": items[0].get("version") if items else None,
    "theme": items[0].get("theme") if items else None,
    "total": len(items),
    "generatedAt": "2026-08-27",
    "groups": [{"key": gk, "label": gl} for gk, gl in GROUP_LABELS.items()],
    "fields": [
        {"key": key, "label": label, "group": group}
        for key, ci, group, label in FIELD_MAP
    ],
    "cat1Order": cat1_order,
    "cat1Cat2": cat1_cat2,
    "stats": {
        "byCat1": [{"name": k, "count": counters["cat1"].get(k, 0)} for k in cat1_order],
        "byType": counters["type"],
        "byCaliber": counters["caliber"],
        "byFreq": counters["freq"],
        "byUnit": counters["unit"],
        "cat1Count": len(cat1_order),
        "cat2Count": len(counters["cat2"]),
    },
}

js_body = json.dumps({"meta": meta, "items": items}, ensure_ascii=False, indent=1)

header = """/* ============================================================
   指标标准库数据（指标查询模块数据源）
   依据《金控指标数据梳理表-指标标准版【清洁版】.xlsx》
   sheet: 银行指标清单 · 共 %(total)d 条指标标准 · 自动生成于 2026-08-27
   结构: INDICATOR_STANDARD = { meta: {...}, items: [...] }
   ============================================================ */
""" % {"total": len(items)}

with open(OUT, "w", encoding="utf-8") as f:
    f.write(header)
    f.write("const INDICATOR_STANDARD = ")
    f.write(js_body)
    f.write(";\n")

print("OK items:", len(items), "cat1:", len(cat1_order))
