# -*- coding: utf-8 -*-
"""探测 金控指标数据梳理表-指标标准版【清洁版】 中的层级格式（加粗/缩进/前导空格）"""
import openpyxl
import json
import os

BASE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(BASE, "金控指标数据梳理表-指标标准版【清洁版】 .xlsx")
OUT = os.path.join(BASE, "hierarchy_clean.json")

wb = openpyxl.load_workbook(SRC)

result = {}
for sn in wb.sheetnames:
    ws = wb[sn]
    special = []
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 220)):
        for cell in row:
            v = cell.value
            if v is None:
                continue
            raw = str(v)
            lead = len(raw) - len(raw.lstrip())
            bold = bool(cell.font.bold) if cell.font and cell.font.bold is not None else False
            indent = float(cell.alignment.indent) if cell.alignment and cell.alignment.indent else 0
            if bold or indent or lead:
                special.append({
                    "row": cell.row, "col": cell.column,
                    "name": raw.strip(), "bold": bold,
                    "indent": indent, "lead": lead
                })
    result[sn] = {"max_row": ws.max_row, "max_col": ws.max_column, "special": special}

with open(OUT, "w", encoding="utf-8") as f:
    json.dump(result, f, ensure_ascii=False, indent=1)

for sn, d in result.items():
    print(sn, "max_row", d["max_row"], "特殊格式格数", len(d["special"]))
