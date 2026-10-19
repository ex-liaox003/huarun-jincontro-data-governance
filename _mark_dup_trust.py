# -*- coding: utf-8 -*-
"""标记信托报表.xlsx《合并》sheet 中的重复指标：B列=是，C列=重复名称。另存新文件。"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl

SRC = '信托报表.xlsx'
DST = '信托报表-合并-重复标记.xlsx'
SHEET = '合并'
DATA_START = 3
COL_B, COL_C, COL_F = 2, 3, 6   # 是否重复 / 重复名称 / 指标名称

wb = openpyxl.load_workbook(SRC)
ws = wb[SHEET]

groups = {}
for r in range(DATA_START, ws.max_row + 1):
    name = str(ws.cell(r, COL_F).value or '')
    for ch in (' ', '　', '\n', '\r'):
        name = name.replace(ch, '')
    if not name:
        continue
    groups.setdefault(name, []).append(r)

dup_groups = {n: rs for n, rs in groups.items() if len(rs) > 1}

for name, rs in dup_groups.items():
    for r in rs:
        ws.cell(r, COL_B).value = '是'
        ws.cell(r, COL_C).value = name

wb.save(DST)
print(f'指标总行数: {sum(len(v) for v in groups.values())}')
print(f'重复组数: {len(dup_groups)}')
print(f'标是行数: {sum(len(v) for v in dup_groups.values())}')
for n, rs in sorted(dup_groups.items(), key=lambda kv: kv[1][0]):
    print(f'  {n}: 行{rs}')
