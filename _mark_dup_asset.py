# -*- coding: utf-8 -*-
"""标记《资产指标清单》sheet 中的重复指标：B列=是，C列=重复名称。另存新文件。"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl

SRC = '金控&资产指标数据梳理表.xlsx'
DST = '金控&资产指标数据梳理表-重复标记.xlsx'
SHEET = '资产指标清单'
DATA_START = 5           # 行1-3表头，行4=示例，数据从第5行开始
COL_B, COL_C, COL_J = 2, 3, 10   # 是否重复 / 重复名称 / 金控指标名称

wb = openpyxl.load_workbook(SRC)
ws = wb[SHEET]

# 1. 名称 -> [行号]
groups = {}
for r in range(DATA_START, ws.max_row + 1):
    name = str(ws.cell(r, COL_J).value or '').replace(' ', '').replace('　', '').replace(chr(10), '').replace(chr(13), '')
    if not name or name == '个人AUM月日均' and str(ws.cell(r, 1).value).strip() == '示例':
        continue
    groups.setdefault(name, []).append(r)

dup_groups = {n: rs for n, rs in groups.items() if len(rs) > 1}

# 2. 标记（原表3行"关注资产率"C列误填了"否"，一并覆写为正确名称）
overwritten = []
for name, rs in dup_groups.items():
    for r in rs:
        ws.cell(r, COL_B).value = '是'
        cur = str(ws.cell(r, COL_C).value or '').strip()
        if cur and cur != '否':
            overwritten.append((r, cur))
        else:
            ws.cell(r, COL_C).value = name

# 3. 检查：原有"是"但不在任何同名组的行（跨名称别名），保留
pre_yes = {r for r in range(DATA_START, ws.max_row + 1)
           if str(ws.cell(r, COL_B).value or '').strip() == '是'}
group_rows = {r for rs in dup_groups.values() for r in rs}
kept_cross = sorted(pre_yes - group_rows)

wb.save(DST)
print(f'同名重复组: {len(dup_groups)}')
print(f'标是总行数: {len(group_rows)}')
print(f'保留原有C值(非空非否): {overwritten}')
print(f'原"是"但不在同名组: {kept_cross}')
for n, rs in sorted(dup_groups.items(), key=lambda kv: kv[1][0]):
    print(f'  {n}: 行{rs}')
